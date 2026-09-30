"""Generate DMS-Raptor User Manual v1.7.0 using python-docx."""
import os
import sys
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE

OUT_PATH = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "DMS-Raptor_User_Manual_v1.7.0.docx",
)


def set_cell(cell, text, bold=False, size=10, color=None):
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Header
    for i, h in enumerate(headers):
        set_cell(table.rows[0].cells[i], h, bold=True, size=9)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            set_cell(table.rows[ri + 1].cells[ci], str(val), size=9)
    return table


def build():
    doc = Document()

    # --- Styles ---
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)

    # ====================================================================
    # TITLE PAGE
    # ====================================================================
    doc.add_paragraph("")
    doc.add_paragraph("")
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run("DMS-Raptor")
    run.bold = True
    run.font.size = Pt(36)
    run.font.color.rgb = RGBColor(0, 188, 212)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = sub.add_run("Dynamic Model Switching for Real-Time UAV Inspection")
    r2.font.size = Pt(16)
    r2.font.color.rgb = RGBColor(100, 100, 100)

    ver = doc.add_paragraph()
    ver.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = ver.add_run("User Manual \u2014 Version 1.7.0")
    r3.bold = True
    r3.font.size = Pt(14)

    doc.add_paragraph("")
    cr = doc.add_paragraph()
    cr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r4 = cr.add_run("\u00A9 2025\u20132026 Teja Chirravuri. All rights reserved.")
    r4.font.size = Pt(10)
    r4.font.color.rgb = RGBColor(120, 120, 120)

    doc.add_page_break()

    # ====================================================================
    # TABLE OF CONTENTS (manual placeholder)
    # ====================================================================
    doc.add_heading("Table of Contents", level=1)
    toc_items = [
        "1. Introduction",
        "2. System Requirements",
        "3. Installation & Launch",
        "4. Application Overview",
        "5. Tab 1: Setup",
        "6. Tab 2: Monitor",
        "7. Tab 3: Results",
        "8. Tab 4: Batch Run",
        "9. Tab 5: Frame Extraction",
        "10. Tab 6: Validation",
        "11. Tab 7: Thesis Experiments",
        "12. Switching Policies",
        "13. Output Files & Folder Structure",
        "14. Troubleshooting",
        "15. Changelog (v1.6.0 \u2192 v1.7.0)",
    ]
    for item in toc_items:
        p = doc.add_paragraph(item)
        p.paragraph_format.space_after = Pt(2)
    doc.add_page_break()

    # ====================================================================
    # 1. INTRODUCTION
    # ====================================================================
    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph(
        "DMS-Raptor is a desktop GUI application for complexity-aware dynamic model "
        "switching in real-time object detection on UAV (Unmanned Aerial Vehicle) "
        "platforms. It evaluates scene complexity using lightweight proxy metrics "
        "(Laplacian variance, histogram entropy) and automatically selects between "
        "a fast YOLOv8n model and a more accurate YOLOv8s model on a per-frame basis."
    )
    doc.add_paragraph(
        "Version 1.7.0 introduces the Thesis Experiments tab with 8 sub-tabs "
        "producing 18 publication-quality plot types, an integrated DOCX report "
        "generator, mouse scroll-wheel zoom and click-drag panning on latency "
        "subplots, and per-policy switching signal display."
    )

    # ====================================================================
    # 2. SYSTEM REQUIREMENTS
    # ====================================================================
    doc.add_heading("2. System Requirements", level=1)
    add_table(doc,
        ["Component", "Requirement"],
        [
            ["Python", "3.9 or later"],
            ["OS", "Windows 10/11, Linux (Ubuntu 20.04+)"],
            ["GPU (optional)", "NVIDIA CUDA-capable GPU (recommended)"],
            ["RAM", "8 GB minimum, 16 GB recommended"],
            ["Libraries", "PyQt5, ultralytics, numpy, matplotlib, opencv-python"],
            ["Conda env", "uav-thesis (recommended)"],
        ],
    )

    # ====================================================================
    # 3. INSTALLATION & LAUNCH
    # ====================================================================
    doc.add_heading("3. Installation & Launch", level=1)
    doc.add_heading("3.1 Install Dependencies", level=2)
    doc.add_paragraph(
        "Activate your conda environment and install the required packages:"
    )
    code = doc.add_paragraph()
    r = code.add_run(
        "conda activate uav-thesis\n"
        "pip install PyQt5 ultralytics matplotlib opencv-python numpy"
    )
    r.font.name = "Consolas"
    r.font.size = Pt(10)

    doc.add_heading("3.2 Launch the Application", level=2)
    doc.add_paragraph("From the project root directory:")
    code2 = doc.add_paragraph()
    r2 = code2.add_run("python main.py")
    r2.font.name = "Consolas"
    r2.font.size = Pt(10)
    doc.add_paragraph(
        "The DMS-Raptor window will open with 7 tabs. The application icon "
        "is now visible in the Windows taskbar."
    )

    # ====================================================================
    # 4. APPLICATION OVERVIEW
    # ====================================================================
    doc.add_heading("4. Application Overview", level=1)
    doc.add_paragraph(
        "DMS-Raptor v1.7.0 has 7 main tabs, each serving a distinct purpose:"
    )
    add_table(doc,
        ["Tab", "Purpose"],
        [
            ["Setup", "Configure policies, models, parameters, and input source"],
            ["Monitor", "Live video feed with detection overlay and rolling stats"],
            ["Results", "Summary table, quick plots (subplots + zoom), data explorer"],
            ["Batch Run", "One-click batch processing across multiple dataset subfolders"],
            ["Frame Extraction", "Extract frames + YOLO-format labels from both models"],
            ["Validation", "Proxy validation, correlation analysis, proxy size grid search"],
            ["Thesis Experiments", "Full thesis experiment suite with 18 publication-quality plots, EDA, and DOCX report generation"],
        ],
    )
    doc.add_paragraph(
        "The application uses a dark theme by default and stores persistent "
        "run history as JSON files in the history/ directory."
    )

    # ====================================================================
    # 5. TAB 1: SETUP
    # ====================================================================
    doc.add_heading("5. Tab 1: Setup", level=1)
    doc.add_heading("5.1 Input Source", level=2)
    doc.add_paragraph(
        "Select the input mode (Video File, Live Stream, or Image Folder) and "
        "browse to your source. For video files, supported formats include "
        ".mp4, .avi, .mov, .mkv, and .webm."
    )
    doc.add_heading("5.2 Model Paths", level=2)
    doc.add_paragraph(
        "Specify paths to two YOLO model weights:\n"
        "\u2022 Model-n: The lightweight model (e.g., yolov8n_glass.pt)\n"
        "\u2022 Model-s: The heavier model (e.g., yolov8s_glass.pt)\n\n"
        "Optional TensorRT export is available for GPU acceleration."
    )
    doc.add_heading("5.3 Policy Selection", level=2)
    doc.add_paragraph(
        "v1.7.0 offers 7 switching policies. Select one or more policies to "
        "evaluate sequentially:"
    )
    add_table(doc,
        ["Policy", "Description"],
        [
            ["n_only", "Always use YOLOv8n (baseline fast)"],
            ["s_only", "Always use YOLOv8s (baseline accurate)"],
            ["entropy_only", "Switch based on histogram entropy H"],
            ["combined", "Switch using weighted C-score = alpha*L + (1-alpha)*H"],
            ["combined_hyst", "Combined with hysteresis (min dwell, switch limits)"],
            ["conf_ema", "Switch based on EMA-smoothed confidence drop"],
            ["niqe_switch", "Switch based on EMA-smoothed NIQE quality score"],
        ],
    )
    doc.add_paragraph(
        "Note: Three legacy policies (combined_inv, combined_hyst_inv, ema_switch) "
        "have been removed in v1.6.0. Old history files referencing them still load correctly."
    )
    doc.add_heading("5.4 Parameters", level=2)
    doc.add_paragraph(
        "The Setup tab exposes only parameters relevant to the active policies. "
        "Each parameter has an 'Auto' checkbox that uses the default value."
    )
    doc.add_paragraph("Key parameter groups:")
    params = [
        ("Scene Analysis", "alpha (L weight), proxy_size, hist_bins, probe_every_k"),
        ("Thresholds", "c_low, c_high, combined_mid"),
        ("Confidence-EMA", "conf_ema_fast_beta, conf_ema_slow_beta, conf_ema_c_high, conf_ema_c_low"),
        ("NIQE-Switch", "niqe_fast_beta, niqe_slow_beta, niqe_c_high, niqe_c_low"),
        ("Stability", "min_dwell_frames, max_switches_per_100"),
        ("Processing", "conf_min, iou_nms, conf_show, max_frames, stride, latency_smooth_window"),
    ]
    add_table(doc, ["Group", "Parameters"], params)

    doc.add_heading("5.5 Running", level=2)
    doc.add_paragraph(
        "Click the green 'Run Selected Policies' button to start inference. "
        "The app switches to the Monitor tab automatically. Use the red 'Stop' "
        "button to gracefully halt processing at any time."
    )

    # ====================================================================
    # 6. TAB 2: MONITOR
    # ====================================================================
    doc.add_heading("6. Tab 2: Monitor", level=1)
    doc.add_paragraph(
        "The Monitor tab displays a live annotated video feed with detection "
        "bounding boxes, class labels, and confidence scores. A dual overlay "
        "shows both model outputs side by side."
    )
    doc.add_paragraph(
        "Below the video feed, rolling statistics charts show real-time "
        "latency, C-score, model choice, and detection count. The status bar "
        "at the bottom shows frame progress (e.g., 'Frame 234/5000 (5%)')."
    )

    # ====================================================================
    # 7. TAB 3: RESULTS
    # ====================================================================
    doc.add_heading("7. Tab 3: Results", level=1)
    doc.add_heading("7.1 Results Table", level=2)
    doc.add_paragraph(
        "The results table displays all completed runs with 9 columns: "
        "Policy, Mode, Mean T (ms), P95 T (ms), Slow %, Sw/100, Frames, "
        "Source (History/Session), and Video info."
    )
    doc.add_paragraph(
        "History runs (loaded from disk on startup) are marked as 'History' "
        "in the Source column. Current session runs are marked as 'Session'. "
        "This distinction also appears visually in the plots."
    )
    doc.add_heading("7.2 Quick Plots", level=2)
    doc.add_paragraph(
        "The Quick Plots panel contains 4 sub-tabs:"
    )
    plots_info = [
        ("Timing Breakdown", "Stacked bar chart showing T_scene, T_ctrl, T_infer_n, T_infer_s per policy"),
        ("Latency Comparison", "Grouped bar chart comparing Mean, P95, and P99 latencies"),
        ("Latency Subplots", "4 vertically stacked waveform subplots (T_total, T_scene, T_ctrl, T_infer) with zoom controls"),
        ("Switching Signals", "Per-policy switching signal: C-score for scene-analysis policies (entropy_only, combined, combined_hyst), conf_drop for conf_ema, NIQE score for niqe_switch. Not shown for n_only/s_only baselines."),
    ]
    add_table(doc, ["Sub-tab", "Content"], plots_info)

    doc.add_heading("7.3 Latency Subplots & Zoom", level=2)
    doc.add_paragraph(
        "The Latency Subplots view shows T_total, T_scene, T_ctrl, and T_infer "
        "as 4 vertically stacked plots sharing the same x-axis (frame index). "
        "History runs appear as dashed, semi-transparent lines; session runs "
        "appear as solid lines with full opacity."
    )
    doc.add_paragraph(
        "Zoom and pan controls:"
    )
    zoom_controls = [
        ("Mouse scroll-wheel", "Zoom in/out on any subplot"),
        ("Click-and-drag", "Pan through frames when zoomed"),
        ("[-] / [+] buttons", "Step zoom level"),
        ("Zoom slider", "Fine zoom control"),
        ("\"Fit All\" button", "Reset to full view"),
        ("Pan slider", "Appears when zoomed, allows scrolling through data"),
    ]
    add_table(doc, ["Control", "Action"], zoom_controls)

    doc.add_heading("7.4 Switching Signal Display", level=2)
    doc.add_paragraph(
        "In v1.7.0, the Switching Signals sub-tab displays the appropriate "
        "per-policy switching signal for each run. Scene-analysis policies "
        "(entropy_only, combined, combined_hyst) show the C-score with their "
        "respective c_low/c_high threshold traces as dotted lines. The conf_ema "
        "policy shows conf_drop, and niqe_switch shows the NIQE score. Baseline "
        "policies (n_only, s_only) are excluded as they have no switching signal. "
        "This fixes a v1.6.0 issue where the C-score plot showed irrelevant "
        "zeros for conf_ema and niqe_switch policies."
    )
    doc.add_heading("7.5 Data Explorer", level=2)
    doc.add_paragraph(
        "The Data Explorer offers flexible plotting with 4 modes:\n"
        "\u2022 Trace (per-frame): Line, Scatter, or Area plots of any metric vs frame index\n"
        "\u2022 Aggregate Comparison: Bar charts of summary metrics\n"
        "\u2022 Distribution: Histogram, Box Plot, or Violin Plot of per-frame data\n"
        "\u2022 Correlation Heatmap: Pearson correlation matrix of all trace metrics"
    )
    doc.add_heading("7.6 Export", level=2)
    doc.add_paragraph(
        "Use the buttons below the table to:\n"
        "\u2022 Export CSV: Save all results to a CSV file\n"
        "\u2022 Export Plots: Save all quick plot figures as PNGs (220 DPI). "
        "The switching signal plot is exported as switching_signals.png.\n"
        "\u2022 Open Video Folder: Navigate to annotated output videos\n"
        "\u2022 Delete Selected / Clear All: Manage displayed results"
    )

    # ====================================================================
    # 8. TAB 4: BATCH RUN
    # ====================================================================
    doc.add_heading("8. Tab 4: Batch Run", level=1)
    doc.add_paragraph(
        "The Batch Run tab enables one-click processing of multiple datasets "
        "organized in subfolders. This is ideal for running all policies across "
        "glass, composite, porcelain, and other material datasets simultaneously."
    )
    doc.add_heading("8.1 Folder Structure", level=2)
    doc.add_paragraph(
        "Organize your data as follows:"
    )
    code3 = doc.add_paragraph()
    r3 = code3.add_run(
        "source_videos/            models/\n"
        "  glass/                    glass/\n"
        "    video1.mp4                yolov8n_glass.pt\n"
        "    video2.mp4                yolov8s_glass.pt\n"
        "  composite/                composite/\n"
        "    video3.mp4                yolov8n_comp.pt\n"
        "                              yolov8s_comp.pt"
    )
    r3.font.name = "Consolas"
    r3.font.size = Pt(9)

    doc.add_heading("8.2 Workflow", level=2)
    steps = [
        "Set the Source Videos Folder, Models Folder, and Output Folder paths",
        "Click 'Scan Folders' to auto-discover datasets with matching models",
        "Review the Discovered Datasets table (shows subfolder, video count, model availability)",
        "Select desired policies via checkboxes (all 7 checked by default)",
        "Configure options: Save annotated videos, Generate comparison plots, Export CSV, Run Proxy Validation",
        "Set Device (cpu/cuda:0) and Image Size (default 640)",
        "Click 'Run All' to start batch processing",
        "Monitor progress via the progress bar and status label",
        "Click 'Open Output' to view results when complete",
    ]
    for i, step in enumerate(steps, 1):
        doc.add_paragraph(f"{i}. {step}")

    doc.add_heading("8.3 Model Matching", level=2)
    doc.add_paragraph(
        "The scanner matches .pt files to n/s roles by looking for keywords "
        "in the filename: 'nano', '_n.', '_n_', 'yolov8n' map to Model-n; "
        "'small', '_s.', '_s_', 'yolov8s' map to Model-s. If no keywords match, "
        "alphabetical order is used (first = n, second = s)."
    )
    doc.add_heading("8.4 Output Structure", level=2)
    doc.add_paragraph(
        "Annotated video filenames now include the input video name for "
        "clarity. The naming pattern is annotated_{video_stem}_{policy}.mp4."
    )
    code4 = doc.add_paragraph()
    r4 = code4.add_run(
        "output/\n"
        "  glass/\n"
        "    annotated_video1_n_only.mp4\n"
        "    annotated_video1_conf_ema.mp4\n"
        "    comparison_plots.png\n"
        "    summary.csv\n"
        "    validation_report.txt\n"
        "    correlation_matrix.png\n"
        "  porcelain/\n"
        "    annotated_porcelain_maybe_conf_ema.mp4\n"
        "    ...\n"
        "  composite/\n"
        "    ...\n"
        "  batch_summary.csv"
    )
    r4.font.name = "Consolas"
    r4.font.size = Pt(9)

    # ====================================================================
    # 9. TAB 5: FRAME EXTRACTION
    # ====================================================================
    doc.add_heading("9. Tab 5: Frame Extraction", level=1)
    doc.add_paragraph(
        "The Frame Extraction tab extracts image frames and YOLO-format labels "
        "from both the n_only and s_only model runs. This produces a ready-to-use "
        "dataset for future model retraining."
    )
    doc.add_heading("9.1 Configuration", level=2)
    config_items = [
        ("Source Folder", "Folder containing subfolders with videos (same structure as Batch Run)"),
        ("Model-n Path", "Path to the YOLOv8n .pt weights file"),
        ("Model-s Path", "Path to the YOLOv8s .pt weights file"),
        ("Output Folder", "Where to save extracted frames and labels"),
        ("Image Format", "jpg (smaller) or png (lossless)"),
        ("Confidence Threshold", "Minimum confidence for including a detection (default 0.25)"),
        ("Device", "cpu or cuda:N"),
        ("Max Frames", "Limit frames per video (0 = all)"),
        ("Frame Stride", "Process every Nth frame (1 = every frame)"),
        ("Image Size", "YOLO inference resolution (default 640)"),
    ]
    add_table(doc, ["Setting", "Description"], config_items)

    doc.add_heading("9.2 Output Structure", level=2)
    code5 = doc.add_paragraph()
    r5 = code5.add_run(
        "output/\n"
        "  glass/\n"
        "    n_only/\n"
        "      images/\n"
        "        frame_000000.jpg\n"
        "        frame_000001.jpg\n"
        "      labels/\n"
        "        frame_000000.txt\n"
        "        frame_000001.txt\n"
        "    s_only/\n"
        "      images/\n"
        "        frame_000000.jpg\n"
        "      labels/\n"
        "        frame_000000.txt\n"
        "  composite/\n"
        "    n_only/...\n"
        "    s_only/..."
    )
    r5.font.name = "Consolas"
    r5.font.size = Pt(9)

    doc.add_heading("9.3 YOLO Label Format", level=2)
    doc.add_paragraph(
        "Each label .txt file contains one line per detection in standard YOLO format:"
    )
    code6 = doc.add_paragraph()
    r6 = code6.add_run("class_id  x_center  y_center  width  height")
    r6.font.name = "Consolas"
    r6.font.size = Pt(10)
    doc.add_paragraph(
        "All coordinates are normalized to [0, 1] relative to the image dimensions. "
        "This format is directly compatible with YOLO training pipelines."
    )

    # ====================================================================
    # 10. TAB 6: VALIDATION
    # ====================================================================
    doc.add_heading("10. Tab 6: Validation", level=1)
    doc.add_paragraph(
        "The Validation tab runs both YOLOv8n and YOLOv8s on every frame of a "
        "video, then measures whether the C-score (from L and H proxies) "
        "actually predicts when the heavier model was needed. This is offline "
        "validation only and does not affect the deployed switching system."
    )
    doc.add_heading("10.1 Validation Configuration", level=2)
    doc.add_paragraph(
        "Specify the video source, both model paths, device, alpha (L weight), "
        "confidence threshold, IoU match threshold, image size, and max frames."
    )
    doc.add_heading("10.2 Results Tabs", level=2)
    val_tabs = [
        ("Report", "Text-based summary of correlation analysis with Pearson coefficients"),
        ("Correlation Scatter", "Interactive scatter plot of any proxy vs any disagreement metric"),
        ("Time Series Overlay", "C-score, L, H overlaid with disagreement score over time"),
        ("Correlation Matrix", "Heatmap of all proxy-metric correlations"),
        ("Sensitivity Analysis", "Heatmap showing how different alpha values affect correlation"),
        ("Detection Comparison", "Side-by-side detection count comparison between models"),
        ("Proxy Size Study", "Grid search for optimal proxy_size x hist_bins combination"),
    ]
    add_table(doc, ["Sub-tab", "Description"], val_tabs)

    doc.add_heading("10.3 Proxy Size Study", level=2)
    doc.add_paragraph(
        "The Proxy Size Study performs a grid search over proxy_size and hist_bins "
        "combinations to find the optimal configuration for scene complexity analysis."
    )
    doc.add_paragraph(
        "Configuration:\n"
        "\u2022 proxy_size range: min, max, step (default 64 to 640, step 64)\n"
        "\u2022 hist_bins range: min, max, step (default 32 to 256, step 32)\n\n"
        "The search computes the C-score for sampled frames (up to 200) at each "
        "combination and measures |Pearson r| correlation with model disagreement. "
        "Results are displayed as a heatmap with the best combination marked by "
        "a white star."
    )
    doc.add_paragraph(
        "Prerequisites: Run a standard validation first to generate the "
        "disagreement data needed for the grid search."
    )

    # ====================================================================
    # 11. TAB 7: THESIS EXPERIMENTS
    # ====================================================================
    doc.add_heading("11. Tab 7: Thesis Experiments", level=1)

    doc.add_heading("11.1 Overview", level=2)
    doc.add_paragraph(
        "The Thesis Experiments tab is a dedicated experiment dashboard for thesis "
        "work. It provides publication-quality plots, statistical analysis, and a "
        "downloadable DOCX research report. The tab contains 8 sub-tabs covering "
        "latency analysis, switching behavior, proxy signals, model comparison, "
        "hyperparameter sensitivity, exploratory data analysis, feature engineering, "
        "and a comprehensive research summary."
    )

    doc.add_heading("11.2 Control Panel", level=2)
    doc.add_paragraph(
        "The control panel at the top of the Thesis Experiments tab provides "
        "the following controls:"
    )
    control_items = [
        ("Video source selector", "Browse to select the input video file for experiments"),
        ("Model path selector", "Set paths to both YOLOv8n and YOLOv8s model weights"),
        ("Policy filter checkboxes", "Select which switching policies to include in the experiment run"),
        ("Run Full Suite button", "Execute all selected policies and generate all 18 plot types"),
        ("Stop button", "Halt the experiment run gracefully"),
        ("Export All Plots button", "Save all 18 plot types as high-DPI PNG images"),
        ("Download DOCX Report button", "Generate and save a comprehensive 12-chapter Word document"),
        ("Progress bar", "Displays progress during experiment runs"),
    ]
    add_table(doc, ["Control", "Description"], control_items)

    doc.add_heading("11.3 Sub-tabs", level=2)
    doc.add_paragraph(
        "The Thesis Experiments tab contains 8 sub-tabs, each presenting a "
        "different aspect of the experiment results:"
    )
    subtabs = [
        ("Latency Dashboard",
         "CDF plots, box plots, timing breakdown stacked bars, and latency time "
         "series with P95 band. Provides a comprehensive view of per-policy "
         "latency characteristics."),
        ("Switching Analysis",
         "Choice heatmap displaying a binary strip per policy showing model "
         "selection over time, switching rate computed via a rolling window, and "
         "dwell histogram showing run-length distributions for each policy."),
        ("Proxy Signals",
         "Laplacian vs entropy scatter with regression line, detection count "
         "scatter (n vs s detections per frame), and confidence comparison scatter "
         "across policies."),
        ("Model Comparison",
         "Detection count comparison between the YOLOv8n and YOLOv8s models "
         "derived from validation data. Shows per-frame agreement and disagreement "
         "patterns."),
        ("Hyperparameter Sensitivity",
         "Threshold sweep heatmap showing policy behavior across parameter ranges, "
         "and Pareto front plotting switching rate vs s-model usage to identify "
         "optimal operating points."),
        ("EDA\u2014Statistics",
         "Descriptive statistics table with mean, median, std, min, max, P95, P99, "
         "IQR, skew, and kurtosis for all metrics. KDE distribution plots, outlier "
         "detection analysis, and Q-Q plots for normality assessment."),
        ("Feature Engineering",
         "Full correlation matrix heatmap of all recorded features, feature "
         "importance ranking by |Pearson r| vs model choice, decision boundary "
         "visualization (Laplacian vs entropy colored by model choice), and signal "
         "lag cross-correlation analysis."),
        ("Research Summary",
         "Auto-generated HTML summary covering system architecture, per-frame "
         "pipeline description, mathematical foundations, policy descriptions, "
         "experiment results, and key findings. This summary is rendered inline "
         "for quick review."),
    ]
    for name, desc in subtabs:
        doc.add_heading(f"11.3.{subtabs.index((name, desc)) + 1} {name}", level=3)
        doc.add_paragraph(desc)

    doc.add_heading("11.4 DOCX Report", level=2)
    doc.add_paragraph(
        "Click the 'Download DOCX Report' button to generate a comprehensive "
        "12-chapter Word document suitable for inclusion in a thesis or research "
        "publication. The report chapters are:"
    )
    chapters = [
        ("1", "Executive Summary"),
        ("2", "System Architecture"),
        ("3", "Mathematical Foundations"),
        ("4", "Switching Policies (detailed descriptions with formulas)"),
        ("5", "Per-Frame Pipeline Walkthrough"),
        ("6", "Experiment Configuration"),
        ("7", "Policy Results"),
        ("8", "Validation Findings"),
        ("9", "EDA & Statistical Analysis"),
        ("10", "Feature Engineering Analysis"),
        ("11", "Limitations & Future Work"),
        ("12", "Conclusions"),
    ]
    add_table(doc, ["Chapter", "Title"], chapters)
    doc.add_paragraph(
        "The report includes human-readable formulas, algorithm steps, and "
        "detailed discussion of all experiment results. It is saved as a .docx "
        "file that can be further edited in Microsoft Word or LibreOffice Writer."
    )

    doc.add_heading("11.5 Export", level=2)
    doc.add_paragraph(
        "Click the 'Export All Plots' button to save all 18 plot types as "
        "high-DPI PNG images. Each plot is exported at publication quality "
        "suitable for direct inclusion in thesis documents or journal submissions. "
        "The exported files are saved to a user-selected output directory."
    )

    # ====================================================================
    # 12. SWITCHING POLICIES
    # ====================================================================
    doc.add_heading("12. Switching Policies", level=1)
    doc.add_paragraph(
        "DMS-Raptor v1.7.0 supports 7 switching policies. Each policy uses "
        "a different strategy to decide whether to run YOLOv8n (fast) or "
        "YOLOv8s (accurate) on each frame."
    )
    policies = [
        ["n_only", "Always selects YOLOv8n. Serves as the fast baseline \u2014 "
         "lowest latency but may miss detections in complex scenes.", "None"],
        ["s_only", "Always selects YOLOv8s. Serves as the accuracy baseline \u2014 "
         "highest detection quality but slower.", "None"],
        ["entropy_only", "Computes histogram entropy H of the scene proxy. "
         "Switches to s when H exceeds c_high; reverts to n when H drops below c_low.",
         "c_low, c_high, proxy_size, hist_bins"],
        ["combined", "Computes C = alpha * L + (1 - alpha) * H where L is normalized "
         "Laplacian variance and H is normalized entropy. Switches based on c_low/c_high thresholds.",
         "alpha, c_low, c_high, combined_mid, proxy_size, hist_bins"],
        ["combined_hyst", "Same as combined but with hysteresis stabilizers: "
         "min_dwell_frames prevents rapid toggling, max_switches_per_100 limits switch rate.",
         "alpha, c_low, c_high, min_dwell_frames, max_switches_per_100, proxy_size, hist_bins"],
        ["conf_ema", "Monitors detection confidence via dual EMA (fast and slow). "
         "When fast EMA drops significantly below slow EMA (confidence degradation), "
         "switches to s for better accuracy.",
         "conf_ema_fast_beta, conf_ema_slow_beta, conf_ema_c_high, conf_ema_c_low"],
        ["niqe_switch", "Monitors image quality via NIQE score using dual EMA. "
         "When quality degrades (fast EMA rises above slow EMA), switches to s.",
         "niqe_fast_beta, niqe_slow_beta, niqe_c_high, niqe_c_low"],
    ]
    add_table(doc, ["Policy", "Description", "Key Parameters"], policies)

    # ====================================================================
    # 13. OUTPUT FILES
    # ====================================================================
    doc.add_heading("13. Output Files & Folder Structure", level=1)
    doc.add_heading("13.1 History Files", level=2)
    doc.add_paragraph(
        "Each completed run is saved as a JSON file in the history/ directory "
        "at the project root. Files are named with timestamps and loaded "
        "automatically on application startup. Historical runs appear as "
        "'History' in the results table and with dashed lines in plots."
    )
    doc.add_heading("13.2 Annotated Videos", level=2)
    doc.add_paragraph(
        "When 'Save annotated output video' is enabled in Setup (or in Batch Run), "
        "annotated MP4 videos are saved with detection overlays, model choice "
        "indicators, and per-frame statistics. In v1.7.0, annotated video filenames "
        "include the input video name using the pattern "
        "annotated_{video_stem}_{policy}.mp4."
    )
    doc.add_heading("13.3 CSV Exports", level=2)
    doc.add_paragraph(
        "CSV exports include all aggregate metrics (mean/P95/P99 latency, "
        "switch rate, slow percentage) for easy import into spreadsheet "
        "software or plotting tools."
    )

    # ====================================================================
    # 14. TROUBLESHOOTING
    # ====================================================================
    doc.add_heading("14. Troubleshooting", level=1)
    issues = [
        ["Application does not start", "Ensure PyQt5 is installed. Check that you are using the correct conda environment (uav-thesis)."],
        ["Taskbar icon not showing", "Ensure resources/logo.ico exists. The ctypes AppUserModelID call in main.py must execute before QApplication."],
        ["CUDA out of memory", "Reduce Image Size (imgsz) from 640 to 320. Close other GPU-intensive applications."],
        ["TensorRT export fails", "Ensure tensorrt and torch are compatible versions. The app falls back to .pt automatically."],
        ["Models not found in Batch Run", "Ensure model subfolder names match source video subfolder names exactly."],
        ["Slow performance on CPU", "Use stride > 1 to skip frames, or set max_frames to limit processing."],
        ["Old history files show unknown policies", "Legacy policy names (combined_inv, etc.) are handled gracefully. Color fallbacks are used."],
        ["Grid search shows all zeros", "Run a standard validation first. The grid search requires disagreement data from a prior validation run."],
        ["Thesis Experiments tab shows no data", "Run the full experiment suite first using the 'Run Full Suite' button. All sub-tabs populate after the run completes."],
        ["DOCX report generation fails", "Ensure python-docx is installed (pip install python-docx). Check that experiment data is available from a completed run."],
    ]
    add_table(doc, ["Issue", "Solution"], issues)

    # ====================================================================
    # 15. CHANGELOG
    # ====================================================================
    doc.add_heading("15. Changelog (v1.6.0 \u2192 v1.7.0)", level=1)
    changes = [
        "NEW: Thesis Experiments tab with 8 sub-tabs and 18 publication-quality plot types",
        "NEW: Downloadable DOCX thesis experiment report (12 chapters, human-readable formulas)",
        "NEW: Mouse scroll-wheel zoom and click-drag panning on latency subplots",
        "NEW: Per-policy switching signal display (C-score, conf_drop, NIQE as appropriate)",
        "CHANGED: \"C-Score Traces\" sub-tab renamed to \"Switching Signals\"",
        "CHANGED: Annotated video filenames now include input video name",
        "CHANGED: Export filename changed from C_score_traces.png to switching_signals.png",
        "FIXED: C-score plot no longer shows irrelevant zeros for conf_ema/niqe_switch policies",
    ]
    for c in changes:
        doc.add_paragraph(c, style="List Bullet")

    doc.add_page_break()
    # Footer
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run(
        "DMS-Raptor v1.7.0 \u2014 User Manual\n"
        "Generated March 2026\n"
        "\u00A9 2025\u20132026 Teja Chirravuri"
    )
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(150, 150, 150)

    # Save
    doc.save(OUT_PATH)
    print(f"User manual saved to: {OUT_PATH}")


if __name__ == "__main__":
    build()
